"""Excel import (from the project template) and export (same layout)."""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .fields import CONSENT_KEYS, EMP_NO_XLSX, FIELDS, clean

EMP_SHEET = "الموظفون"
REP_SHEET = "المناديب"
MAX_ERRORS = 200
EMP_NO_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9_\-]{0,29}")

_BY_XLSX = {f.xlsx: f for f in FIELDS}


def _header(text):
    return re.sub(r"\s+", " ", str(text or "")).replace("*", "").strip()


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_sheet(ws):
    """-> (list of (row_no, {xlsx_header: cellvalue}), unknown_headers)"""
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [_header(h) for h in next(rows)]
    except StopIteration:
        return [], []
    known = {EMP_NO_XLSX} | set(_BY_XLSX)
    unknown = [h for h in headers if h and h not in known]
    out = []
    for n, row in enumerate(rows, start=2):
        rec = {h: _cell(v) for h, v in zip(headers, row) if h and _cell(v) != ""}
        if rec:
            out.append((n, rec))
    return out, unknown


def import_workbook(stream, actor, dry_run=True):
    wb = load_workbook(stream, read_only=True, data_only=True)
    report = {"created": 0, "updated": 0, "unchanged": 0, "reps": 0, "errors": [],
              "warnings": [], "dry_run": dry_run}
    if EMP_SHEET not in wb.sheetnames:
        report["errors"].append(f"لا توجد ورقة باسم «{EMP_SHEET}» في الملف. استخدم القالب الرسمي.")
        return report

    emp_rows, unknown = _read_sheet(wb[EMP_SHEET])
    rep_rows, unknown2 = _read_sheet(wb[REP_SHEET]) if REP_SHEET in wb.sheetnames else ([], [])
    for h in sorted(set(unknown + unknown2)):
        report["warnings"].append(f"عمود غير معروف تم تجاهله: «{h}»")

    def parse(sheet, row_no, rec):
        emp_no = rec.get(EMP_NO_XLSX, "")
        if not emp_no:
            report["errors"].append(f"{sheet} - صف {row_no}: الرقم الوظيفي فارغ")
            return None
        if not EMP_NO_RE.fullmatch(emp_no):
            report["errors"].append(f"{sheet} - صف {row_no}: الرقم الوظيفي «{emp_no}» غير صالح (حروف إنجليزية وأرقام وشرطات فقط، حتى 30 خانة)")
            return None
        values, bad = {}, []
        for header, raw in rec.items():
            f = _BY_XLSX.get(header)
            if f is None or f.kind == "photo" or f.key in CONSENT_KEYS:
                continue
            value, err = clean(f, raw)
            if err:
                bad.append(f"{f.label}: {err}")
            else:
                values[f.key] = value
        if bad:
            if len(report["errors"]) < MAX_ERRORS:
                report["errors"].append(f"{sheet} - صف {row_no} (موظف {emp_no}): " + " | ".join(bad))
            return None
        return emp_no, values

    seen = set()
    for row_no, rec in emp_rows:
        parsed = parse(EMP_SHEET, row_no, rec)
        if not parsed:
            continue
        emp_no, values = parsed
        if emp_no in seen:
            report["errors"].append(f"{EMP_SHEET} - صف {row_no}: الرقم الوظيفي {emp_no} مكرر في الملف")
            continue
        seen.add(emp_no)
        _upsert(emp_no, values, report, actor, is_rep=None)

    for row_no, rec in rep_rows:
        parsed = parse(REP_SHEET, row_no, rec)
        if not parsed:
            continue
        emp_no, values = parsed
        if db.get_employee(emp_no) is None:
            report["errors"].append(f"{REP_SHEET} - صف {row_no}: الموظف {emp_no} غير موجود في ورقة «{EMP_SHEET}»")
            continue
        _upsert(emp_no, values, report, actor, is_rep=True)
        report["reps"] += 1

    if dry_run:
        db.get_db().rollback()
    return report


def _upsert(emp_no, values, report, actor, is_rep):
    emp = db.get_employee(emp_no)
    if emp is None:
        db.create_employee(emp_no, values, is_rep=bool(is_rep))
        db.audit(actor, "import_create", emp_no, "إنشاء من ملف Excel")
        report["created"] += 1
        return
    changed = [k for k, v in values.items() if emp["data"].get(k) != v]
    flag_change = is_rep and not emp["is_rep"]
    if not changed and not flag_change:
        report["unchanged"] += 1
        return
    data = {**emp["data"], **values}
    cols = {"data": data, "updated_at": db.now_iso()}
    if flag_change:
        cols["is_rep"] = True
    db.update_employee(emp_no, **cols)
    db.audit(actor, "import_update", emp_no, "الحقول: " + "، ".join(changed))
    report["updated"] += 1


# ------------------------------------------------------------------ export
def build_export(employees):
    wb = Workbook()
    ws = wb.active
    ws.title = EMP_SHEET
    _write_sheet(ws, [e for e in employees], rep=False)
    ws2 = wb.create_sheet(REP_SHEET)
    _write_sheet(ws2, [e for e in employees if e["is_rep"]], rep=True)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


STATUS_AR = {"draft": "مسودة", "submitted": "بانتظار المراجعة", "approved": "معتمد", "returned": "معاد للتعديل"}


def _write_sheet(ws, employees, rep):
    cols = [f for f in FIELDS if (f.rep_only if rep else not f.rep_only)]
    headers = [EMP_NO_XLSX + " *"] + [f.xlsx + (" *" if f.required else "") for f in cols]
    if not rep:
        headers += ["حالة الاعتماد", "آخر تحديث"]
    ws.sheet_view.rightToLeft = True
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B2"
    for r, e in enumerate(employees, start=2):
        row = [e["emp_no"]]
        for f in cols:
            if f.kind == "photo":
                ts = e["photos"].get(f.key)
                row.append(("مرفوعة " + db.to_local(ts)[:10]) if ts else "")
            else:
                row.append(e["data"].get(f.key, ""))
        if not rep:
            row += [STATUS_AR.get(e["status"], e["status"]), db.to_local(e["updated_at"])]
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = "@"

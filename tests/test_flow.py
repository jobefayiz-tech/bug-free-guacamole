import io
import os

from openpyxl import load_workbook
from PIL import Image

from app import db
from tests.conftest import CSRF, complete_employee, make_xlsx, png_bytes, post


def import_rows(client, emp_rows, rep_rows=(), dry_run=False, **kw):
    data = {"file": (make_xlsx(emp_rows, rep_rows, **kw), "x.xlsx")}
    if dry_run:
        data["dry_run"] = "1"
    return post(client, "/admin/import", data, content_type="multipart/form-data")


def token_of(app, emp_no):
    with app.app_context():
        return db.get_employee(emp_no)["token"]


def emp(app, emp_no):
    with app.app_context():
        return db.get_employee(emp_no)


# ------------------------------------------------------------------ auth
def test_admin_requires_login(client):
    for url in ("/admin", "/admin/employees", "/admin/export.xlsx", "/admin/import", "/admin/alerts"):
        r = client.get(url)
        assert r.status_code == 302 and "/login" in r.headers["Location"], url


def test_csrf_required(client):
    r = client.post("/login", data={"username": "a", "password": "b"})
    assert r.status_code == 400


def test_login_lockout_and_open_redirect(app, admin, client):
    c2 = app.test_client()
    with c2.session_transaction() as s:
        s["csrf"] = CSRF
    for _ in range(5):
        r = post(c2, "/login", {"username": "boss", "password": "wrong"})
        assert r.status_code == 200
    r = post(c2, "/login", {"username": "boss", "password": "correct-horse-battery"})
    assert r.status_code == 200 and "محاولات كثيرة" in r.get_data(as_text=True)


def test_login_next_cannot_leave_site(app, client):
    from werkzeug.security import generate_password_hash
    with app.app_context():
        db.get_db().execute("INSERT INTO users(username,password_hash,role,created_at) VALUES('hr1',?, 'hr', ?)",
                            (generate_password_hash("correct-horse-battery"), db.now_iso()))
        db.get_db().commit()
    for nxt in ("//evil.com", "/\\evil.com", "https://evil.com"):
        c = app.test_client()
        with c.session_transaction() as s:
            s["csrf"] = CSRF
        r = post(c, "/login?next=" + nxt, {"username": "hr1", "password": "correct-horse-battery"})
        assert r.status_code == 302 and "evil" not in r.headers["Location"], nxt


def test_hr_role_cannot_manage_users(app, client):
    from werkzeug.security import generate_password_hash
    with app.app_context():
        db.get_db().execute("INSERT INTO users(username,password_hash,role,created_at) VALUES('hr1',?, 'hr', ?)",
                            (generate_password_hash("correct-horse-battery"), db.now_iso()))
        db.get_db().commit()
    post(client, "/login", {"username": "hr1", "password": "correct-horse-battery"})
    assert client.get("/admin/users").status_code == 403


# ------------------------------------------------------------------ import
def test_import_dry_run_then_real(app, admin):
    rows = [complete_employee("E1"), complete_employee("E2", name_ar="سعد", id_number="2111111111", mobile="0566666666")]
    r = import_rows(admin, rows, dry_run=True)
    assert "نتيجة التجربة" in r.get_data(as_text=True)
    with app.app_context():
        assert db.all_employees() == []
    r = import_rows(admin, rows)
    assert r.status_code == 200
    assert {e["emp_no"] for e in emp_all(app)} == {"E1", "E2"}
    # idempotent
    import_rows(admin, rows)
    assert len(emp_all(app)) == 2


def emp_all(app):
    with app.app_context():
        return db.all_employees()


def test_import_reports_row_errors_and_skips_only_bad_rows(app, admin):
    rows = [complete_employee("E1"),
            complete_employee("E2", mobile="123"),
            complete_employee("E3", id_number="999"),
            complete_employee("../x"),
            complete_employee("E1", name_ar="مكرر")]
    r = import_rows(admin, rows)
    html = r.get_data(as_text=True)
    assert [e["emp_no"] for e in emp_all(app)] == ["E1"]
    assert "رقم الجوال" in html and "غير صالح" in html and "مكرر" in html


def test_import_updates_without_blanking_and_marks_reps(app, admin):
    import_rows(admin, [complete_employee("E1")])
    import_rows(admin, [dict(emp_no="E1", city="جدة")],
                rep_rows=[dict(emp_no="E1", rep_code="R-1", region="الغربية", supervisor_emp_no="E9", app_name="نينجا", app_id="77")])
    e = emp(app, "E1")
    assert e["data"]["city"] == "جدة" and e["data"]["name_ar"] == "أحمد محمد"  # untouched field kept
    assert e["is_rep"] and e["data"]["app_id"] == "77"


def test_import_rejects_rep_row_for_unknown_employee(app, admin):
    r = import_rows(admin, [complete_employee("E1")], rep_rows=[dict(emp_no="NOPE", rep_code="R")])
    assert "غير موجود" in r.get_data(as_text=True)


def test_import_ignores_consent_columns(app, admin):
    import_rows(admin, [complete_employee("E1", consent="نعم")])
    assert "consent" not in emp(app, "E1")["data"]


def test_import_garbage_file(admin):
    r = post(admin, "/admin/import", {"file": (io.BytesIO(b"not an excel"), "x.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 200 and "تعذّر قراءة الملف" in r.get_data(as_text=True)


# ------------------------------------------------------------------ employee portal
def test_unknown_token_404(client):
    assert client.get("/e/nope").status_code == 404


def test_verify_last4_and_lockout(app, admin):
    import_rows(admin, [complete_employee("E1")])
    tok = token_of(app, "E1")
    c = app.test_client()
    with c.session_transaction() as s:
        s["csrf"] = CSRF
    assert c.get(f"/e/{tok}/form").status_code == 302          # not verified yet
    for _ in range(4):
        r = post(c, f"/e/{tok}", {"last4": "0000"})
        assert "غير مطابقة" in r.get_data(as_text=True)
    r = post(c, f"/e/{tok}", {"last4": "0000"})                  # 5th failure locks
    r = post(c, f"/e/{tok}", {"last4": "7890"})
    assert "تجاوزت عدد المحاولات" in r.get_data(as_text=True)


def test_verify_accepts_mobile_last4_when_the_id_on_file_is_wrong(app, admin):
    import_rows(admin, [complete_employee("E1")])
    tok = token_of(app, "E1")
    c = app.test_client()
    with c.session_transaction() as s:
        s["csrf"] = CSRF
    r = post(c, f"/e/{tok}", {"last4": "567"})                   # too short never matches
    assert "غير مطابقة" in r.get_data(as_text=True)
    r = post(c, f"/e/{tok}", {"last4": "5678"})                  # last 4 of the mobile 0512345678
    assert r.status_code == 302


def verified_client(app, tok, last4="7890"):
    c = app.test_client()
    with c.session_transaction() as s:
        s["csrf"] = CSRF
    r = post(c, f"/e/{tok}", {"last4": last4})
    assert r.status_code == 302, r.get_data(as_text=True)[:300]
    return c


def employee_form_values(**over):
    v = {"action": "submit", "name_ar": "أحمد محمد", "name_en": "Ahmed Mohammed", "nationality": "سعودي", "gender": "ذكر", "birth_date": "1990-05-01",
         "id_type": "هوية وطنية", "id_number": "1234567890", "id_expiry": "2030-01-01", "mobile": "٠٥١٢٣٤٥٦٧٨",
         "city": "الرياض", "emergency_name": "خالد", "emergency_relation": "أخ/أخت", "emergency_phone": "0555555555",
         "consent": "1"}
    v.update(over)
    return v


def test_staff_full_cycle(app, admin):
    import_rows(admin, [complete_employee("E1", mobile="0500000000")])
    c = verified_client(app, token_of(app, "E1"))
    tok = token_of(app, "E1")

    # draft with a bad value: valid fields are saved, bad one is reported and not stored
    r = post(c, f"/e/{tok}/form", employee_form_values(action="draft", mobile="12", city="جدة"))
    html = r.get_data(as_text=True)
    assert "رقم الجوال يجب أن يكون" in html
    e = emp(app, "E1")
    assert e["data"]["city"] == "جدة" and e["data"]["mobile"] == "0500000000" and e["status"] == "draft"

    # submit without consent -> blocked
    r = post(c, f"/e/{tok}/form", employee_form_values(consent=""))
    assert emp(app, "E1")["status"] == "draft" and "هذا الحقل مطلوب" in r.get_data(as_text=True)

    # proper submit
    r = post(c, f"/e/{tok}/form", employee_form_values())
    assert r.status_code == 302 and r.headers["Location"].endswith("/done")
    e = emp(app, "E1")
    assert e["status"] == "submitted" and e["data"]["mobile"] == "0512345678" and e["consent_at"]

    # HR approves
    assert post(admin, "/admin/employees/E1/review", {"action": "approve"}).status_code == 302
    assert emp(app, "E1")["status"] == "approved"

    # employee edits after approval -> goes back for review
    r = post(c, f"/e/{tok}/form", employee_form_values(city="الدمام"))
    assert emp(app, "E1")["status"] == "submitted"

    # HR returns it with a note (required)
    r = post(admin, "/admin/employees/E1/review", {"action": "return", "note": ""})
    assert emp(app, "E1")["status"] == "submitted"
    post(admin, "/admin/employees/E1/review", {"action": "return", "note": "صحح العنوان"})
    e = emp(app, "E1")
    assert e["status"] == "returned"
    assert "صحح العنوان" in c.get(f"/e/{tok}/form").get_data(as_text=True)


def test_cannot_approve_incomplete_record(app, admin):
    import_rows(admin, [dict(emp_no="E5", name_ar="ناقص", id_number="1234567890")])
    c = verified_client(app, token_of(app, "E5"))
    tok = token_of(app, "E5")
    post(c, f"/e/{tok}/form", employee_form_values())   # employee part complete; HR-owned fields still empty
    assert emp(app, "E5")["status"] == "submitted"
    post(admin, "/admin/employees/E5/review", {"action": "approve"})
    assert emp(app, "E5")["status"] == "submitted"       # blocked: job/department missing


def rep_setup(admin, app):
    import_rows(admin, [complete_employee("R1")],
                rep_rows=[dict(emp_no="R1", rep_code="R-1", region="الرياض", supervisor_emp_no="E9")])
    tok = token_of(app, "R1")
    return tok, verified_client(app, tok)


def test_rep_needs_photos_when_has_vehicle(app, admin):
    tok, c = rep_setup(admin, app)
    vals = employee_form_values(app_name="نينجا", app_id="998877", location_consent="نعم",
                                vehicle_type="سيارة الشركة", plate_no="أ ب ج 1234")
    r = post(c, f"/e/{tok}/form", vals)
    html = r.get_data(as_text=True)
    assert emp(app, "R1")["status"] == "draft" and "هذا الحقل مطلوب" in html


def test_rep_full_submit_with_photos_and_privacy(app, admin):
    tok, c = rep_setup(admin, app)
    vals = employee_form_values(app_name="نينجا", app_id="998877", location_consent="نعم",
                                vehicle_type="سيارة الشركة", plate_no="أ ب ج 1234", license_no="1020304050",
                                license_expiry="2029-05-05", registration_expiry="2027-05-05")
    for k in ("front", "back", "right", "left", "odometer"):
        vals[f"photo_{k}"] = (io.BytesIO(png_bytes()), f"{k}.jpg")
    r = post(c, f"/e/{tok}/form", vals, content_type="multipart/form-data")
    assert r.status_code == 302, r.get_data(as_text=True)[:500]
    e = emp(app, "R1")
    assert e["status"] == "submitted" and len(e["photos"]) == 5

    path = os.path.join(app.config["UPLOAD_DIR"], "R1", "photo_front.jpg")
    img = Image.open(path)
    assert max(img.size) <= 1600 and not img.getexif()           # resized, EXIF stripped

    # employee can fetch own photo; anonymous and other employees cannot
    assert c.get(f"/e/{tok}/photo/photo_front").status_code == 200
    anon = app.test_client()
    assert anon.get("/admin/photo/R1/photo_front").status_code == 302      # needs HR login
    assert anon.get(f"/e/{tok}/photo/photo_front").status_code == 404      # link alone is not enough: not verified
    assert admin.get("/admin/photo/R1/photo_front").status_code == 200

    # HR sees the record, sensitive numbers masked until revealed (and reveal is audited)
    page = admin.get("/admin/employees/R1").get_data(as_text=True)
    assert "1234567890" not in page and "7890" in page
    page = admin.get("/admin/employees/R1?reveal=1").get_data(as_text=True)
    assert "1234567890" in page
    with app.app_context():
        acts = [r["action"] for r in db.get_db().execute("SELECT action FROM audit WHERE emp_no='R1'")]
    assert "reveal_sensitive" in acts and "employee_submit" in acts


def test_bad_image_rejected(app, admin):
    tok, c = rep_setup(admin, app)
    vals = employee_form_values(action="draft", photo_front=(io.BytesIO(b"<script>alert(1)</script>"), "evil.jpg"))
    r = post(c, f"/e/{tok}/form", vals, content_type="multipart/form-data")
    assert "الملف ليس صورة صالحة" in r.get_data(as_text=True)
    assert "photo_front" not in emp(app, "R1")["photos"]


def test_employee_cannot_set_hr_owned_fields(app, admin):
    import_rows(admin, [complete_employee("E1")])
    tok = token_of(app, "E1")
    c = verified_client(app, tok)
    post(c, f"/e/{tok}/form", employee_form_values(job_title="مدير عام", department="الإدارة العليا", employment_status="في إجازة"))
    d = emp(app, "E1")["data"]
    assert d["job_title"] == "محاسب" and d["department"] == "المالية"


def test_regenerate_link_invalidates_old(app, admin):
    import_rows(admin, [complete_employee("E1")])
    old = token_of(app, "E1")
    post(admin, "/admin/employees/E1/regen-link")
    assert token_of(app, "E1") != old
    assert app.test_client().get(f"/e/{old}").status_code == 404


# ------------------------------------------------------------------ HR tools
def test_alerts_and_dashboard_and_lists(app, admin):
    import_rows(admin, [complete_employee("E1", id_expiry="2020-01-01"),
                        complete_employee("E2", id_expiry="2099-01-01", id_number="1111111111", mobile="0577777777")])
    html = admin.get("/admin/alerts?days=30").get_data(as_text=True)
    assert "E1" in html and "E2" not in html and "منتهية منذ" in html
    assert admin.get("/admin").status_code == 200
    assert "E1" in admin.get("/admin/employees?q=أحمد").get_data(as_text=True)
    assert "E1" not in admin.get("/admin/employees?status=approved").get_data(as_text=True)
    assert admin.get("/admin/audit").status_code == 200


def test_links_and_whatsapp(app, admin):
    import_rows(admin, [complete_employee("E1")])
    html = admin.get("/admin/links").get_data(as_text=True)
    tok = token_of(app, "E1")
    assert "wa.me/966512345678" in html
    csv_text = admin.get("/admin/links.csv").get_data(as_text=True)
    assert f"/e/{tok}" in csv_text


def test_export_roundtrip(app, admin):
    import_rows(admin, [complete_employee("E1")],
                rep_rows=[dict(emp_no="E1", rep_code="R-1", region="الرياض", supervisor_emp_no="E9", app_name="نينجا", app_id="5")])
    r = admin.get("/admin/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb["الموظفون"].max_row == 2 and wb["المناديب"].max_row == 2
    # the export can be re-imported as is
    r2 = import_rows(admin, [], dry_run=True)  # empty file still parses
    assert r2.status_code == 200
    r3 = post(admin, "/admin/import", {"file": (io.BytesIO(r.data), "export.xlsx"), "dry_run": "1"}, content_type="multipart/form-data")
    html = r3.get_data(as_text=True)
    assert "بدون تغيير" in html and "الأخطاء" not in html


def test_security_headers(client):
    r = client.get("/login")
    assert r.headers["X-Frame-Options"] == "DENY" and r.headers["Referrer-Policy"] == "no-referrer"
    assert "no-store" in r.headers["Cache-Control"]
    assert "style=" not in r.get_data(as_text=True)


def test_hr_edit_cannot_change_consent(app, admin):
    import_rows(admin, [complete_employee("E1")])
    post(admin, "/admin/employees/E1/edit", {"name_ar": "اسم جديد", "consent": "1", "location_consent": "نعم", "city": "أبها"})
    e = emp(app, "E1")
    assert e["data"]["city"] == "أبها" and not e["data"].get("consent")


# ------------------------------------------------------------------ multilingual / non-Arab workers
def test_portal_languages_and_ltr_rtl(app, admin):
    import_rows(admin, [complete_employee("E1", name_en="Ravi Kumar", name_ar="", nationality="هندي")])
    tok = token_of(app, "E1")
    c = verified_client(app, tok)
    for lang, direction, needle in (("en", "ltr", "Full name (English"), ("ur", "rtl", "پورا نام"),
                                    ("bn", "ltr", "পূর্ণ নাম"), ("hi", "ltr", "पूरा नाम"), ("ar", "rtl", "الاسم الكامل")):
        html = c.get(f"/e/{tok}/form?lang={lang}").get_data(as_text=True)
        assert f'lang="{lang}" dir="{direction}"' in html, lang
        assert needle in html, lang
    assert "Ravi Kumar" in c.get(f"/e/{tok}/form?lang=en").get_data(as_text=True)   # shown instead of the empty Arabic name


def test_errors_are_translated_and_values_stay_canonical(app, admin):
    import_rows(admin, [complete_employee("E1", nationality="باكستاني")])
    tok = token_of(app, "E1")
    c = verified_client(app, tok)
    r = post(c, f"/e/{tok}/form?lang=ur", employee_form_values(action="draft", mobile="12", gender="مرد"))
    html = r.get_data(as_text=True)
    assert "موبائل نمبر 05xxxxxxxx" in html            # Urdu error
    assert "فہرست میں سے منتخب کریں" in html             # a select given a translated (non-canonical) value is rejected
    r = post(c, f"/e/{tok}/form?lang=bn", employee_form_values(action="draft", gender="ذكر"))
    assert emp(app, "E1")["data"]["gender"] == "ذكر"      # stored canonical Arabic whatever the UI language


def test_non_arab_worker_can_submit_with_english_name_and_international_emergency_phone(app, admin):
    import_rows(admin, [complete_employee("B1", name_ar="", name_en="Md Rahim", nationality="بنغلاديشي",
                                          id_type="إقامة", id_number="2123456789")])
    tok = token_of(app, "B1")
    c = verified_client(app, tok, last4="6789")
    vals = employee_form_values(name_ar="", name_en="Md Rahim", nationality="بنغلاديشي", id_type="إقامة",
                                id_number="2123456789", iqama_profession="Delivery driver",
                                emergency_phone="+880 1712-345678", emergency_name="Karim", lang="bn")
    r = post(c, f"/e/{tok}/form?lang=bn", vals)
    assert r.status_code == 302, r.get_data(as_text=True)[:400]
    d = emp(app, "B1")["data"]
    assert d["emergency_phone"] == "+8801712345678" and d["name_en"] == "Md Rahim"
    assert emp(app, "B1")["status"] == "submitted"


def test_iqama_profession_required_only_for_residents(app, admin):
    import_rows(admin, [complete_employee("B2", name_en="Ali", id_type="إقامة", id_number="2123456780")])
    tok = token_of(app, "B2")
    c = verified_client(app, tok, last4="6780")
    r = post(c, f"/e/{tok}/form", employee_form_values(id_type="إقامة", id_number="2123456780"))
    assert emp(app, "B2")["status"] == "draft"             # no profession given
    assert "هذا الحقل مطلوب" in r.get_data(as_text=True)


def test_links_pick_language_by_nationality(app, admin):
    import_rows(admin, [complete_employee("E1", nationality="بنغلاديشي", name_en="Rahim"),
                        complete_employee("E2", nationality="باكستاني", name_en="Imran", id_number="1111111111", mobile="0577777777"),
                        complete_employee("E3", nationality="هندي", name_en="Ravi", id_number="1222222222", mobile="0588888888"),
                        complete_employee("E4", nationality="مصري", id_number="1333333333", mobile="0599999999")])
    html = admin.get("/admin/links").get_data(as_text=True)
    for lang_name in ("বাংলা", "اردو", "English", "العربية"):
        assert lang_name in html
    t1, t2, t3, t4 = (token_of(app, f"E{i}") for i in (1, 2, 3, 4))
    csv_text = admin.get("/admin/links.csv").get_data(as_text=True)
    assert f"/e/{t1}?lang=bn" in csv_text and f"/e/{t2}?lang=ur" in csv_text
    assert f"/e/{t3}?lang=en" in csv_text and f"/e/{t4}" in csv_text and f"/e/{t4}?lang" not in csv_text
    # forcing a language and filtering by nationality
    only_hi = admin.get("/admin/links.csv?lang=hi&q=باكستاني").get_data(as_text=True)
    assert f"/e/{t2}?lang=hi" in only_hi and t1 not in only_hi


def test_apps_dropdown_and_nationality_suggestions(app, admin):
    import_rows(admin, [complete_employee("R9")], rep_rows=[dict(emp_no="R9", rep_code="R-9", region="الرياض", supervisor_emp_no="E1")])
    tok = token_of(app, "R9")
    c = verified_client(app, tok)
    html = c.get(f"/e/{tok}/form?lang=en").get_data(as_text=True)
    for name in ("Keeta Food", "Keeta Market", "Ninja", "ToYou", "HungerStation", "Mrsool"):
        assert name in html, name
    assert '<datalist id="dl_nationality">' in html and "Bangladeshi" in html
    r = import_rows(admin, [], rep_rows=[dict(emp_no="R9", app_name="أوبر")])
    assert "اختر من القائمة" in r.get_data(as_text=True)    # an app outside the six is rejected on import

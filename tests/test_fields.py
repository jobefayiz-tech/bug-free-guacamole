import os
import re

from openpyxl import load_workbook

from app.fields import (BY_KEY, EMP_NO_XLSX, FIELDS, OPTIONS, clean, iban_ok, mask, missing, normalize)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def test_arabic_digits_normalized():
    assert normalize("٠٥١٢٣٤٥٦٧٨") == "0512345678"
    assert clean(BY_KEY["mobile"], "٠٥١٢٣٤٥٦٧٨") == ("0512345678", None)


def test_mobile_formats():
    f = BY_KEY["mobile"]
    for raw in ("0512345678", "512345678", "+966512345678", "966 51 234 5678", "00966512345678"):
        assert clean(f, raw) == ("0512345678", None), raw
    for raw in ("0412345678", "05123", "abc"):
        assert clean(f, raw)[1], raw


def test_id_number_rules():
    f = BY_KEY["id_number"]
    assert clean(f, "1234567890")[1] is None
    assert clean(f, "2234567890")[1] is None
    assert clean(f, "3234567890")[1]
    assert clean(f, "123456789")[1]


def test_iban():
    assert iban_ok("SA0380000000608010167519")          # standard published example
    assert not iban_ok("SA0380000000608010167518")      # checksum broken
    assert not iban_ok("GB82WEST12345698765432")
    assert clean(BY_KEY["iban"], "sa03 8000 0000 6080 1016 7519") == ("SA0380000000608010167519", None)


def test_dates():
    f = BY_KEY["id_expiry"]
    assert clean(f, "2030-1-5") == ("2030-01-05", None)  # non zero-padded input is accepted and normalized
    assert clean(f, "15/03/2030") == ("2030-03-15", None)
    assert clean(f, "2030-03-15") == ("2030-03-15", None)
    assert clean(f, "not a date")[1]
    assert clean(BY_KEY["birth_date"], "2024-01-01")[1]  # too young


def test_select_rejects_unknown():
    assert clean(BY_KEY["gender"], "ذكر")[1] is None
    assert clean(BY_KEY["gender"], "x")[1]


def test_mask():
    assert mask(BY_KEY["id_number"], "1234567890") == "••••••7890"
    assert mask(BY_KEY["name_ar"], "أحمد") == "أحمد"
    assert mask(BY_KEY["birth_date"], "1990-05-01").startswith("1990")


def test_missing_vehicle_rules():
    base = {"vehicle_type": "بدون مركبة"}
    keys = missing(base, {}, True, owner="employee")
    assert "plate_no" not in keys and "photo_front" not in keys
    keys = missing({"vehicle_type": "سيارة الشركة"}, {}, True, owner="employee")
    for k in ("plate_no", "license_type", "license_no", "license_expiry", "registration_expiry", "photo_front", "photo_odometer"):
        assert k in keys
    assert "app_id" not in missing({}, {}, False)  # rep-only fields never required for staff


def test_motorcycle_riders_and_licence_type():
    assert clean(BY_KEY["license_type"], "نقل خفيف") == ("نقل خفيف", None)
    assert clean(BY_KEY["license_type"], "دراجة نارية") == ("دراجة نارية", None)
    assert clean(BY_KEY["license_type"], "خاصة")[1]                       # only the listed licence types
    for v in ("دراجة نارية للشركة", "دراجة نارية شخصية"):
        assert clean(BY_KEY["vehicle_type"], v) == (v, None)
        keys = missing({"vehicle_type": v}, {}, True, owner="employee")   # a motorcycle needs the same plate/licence/photos
        for k in ("license_type", "license_no", "plate_no", "photo_front", "photo_odometer"):
            assert k in keys


def test_representative_form_is_the_short_one():
    from app.fields import fields_for
    rep = {f.key for f in fields_for(True)}
    staff = {f.key for f in fields_for(False)}
    full = {f.key for f in fields_for(True, full=True)}
    for k in ("id_number", "id_expiry", "iqama_profession", "vehicle_type", "plate_no", "photo_odometer", "app_name", "app_id"):
        assert k in rep, k
    for k in ("city", "emergency_name", "emergency_phone", "bank_name", "iban", "birth_date", "gender", "email", "passport_no"):
        assert k not in rep and k in staff and k in full, k       # employee file only; HR still sees all of it for a rep
    assert "app_name" not in staff                                   # rep-only fields never reach staff
    assert "city" not in missing({}, {}, True)                       # and a rep is never asked for them


def test_template_matches_fields():
    """tools/make_template.py and app/fields.py must describe the same columns and lists."""
    ns = {"__file__": os.path.join(ROOT, "tools", "make_template.py")}
    src = open(os.path.join(ROOT, "tools", "make_template.py"), encoding="utf-8").read()
    src = src.split("wb = Workbook()")[0]  # definitions only
    exec(compile(src, "make_template", "exec"), ns)
    tpl_emp = [c[0] for c in ns["EMP_COLS"]]
    tpl_rep = [c[0] for c in ns["REP_COLS"]]
    app_emp = [EMP_NO_XLSX] + [f.xlsx for f in FIELDS if not f.rep_only]
    app_rep = [EMP_NO_XLSX] + [f.xlsx for f in FIELDS if f.rep_only]
    assert tpl_emp == app_emp          # same order too: HR pastes the template columns straight into the Google Sheet
    assert tpl_rep == app_rep
    tpl_lists = {tuple(v) for v in ns["LISTS"].values()}
    app_lists = {tuple(v) for v in OPTIONS.values()}
    assert tpl_lists == app_lists
    # required flags agree
    tpl_req = {c[0]: c[1] for c in ns["EMP_COLS"] + ns["REP_COLS"]}
    for f in FIELDS:
        if f.vehicle or f.key == "iqama_profession":
            continue
        assert tpl_req[f.xlsx] == f.required, f.key


def test_shipped_template_has_all_headers():
    wb = load_workbook(os.path.join(ROOT, "employee-data-template.xlsx"))
    heads = {re.sub(r"\s*\*$", "", c.value) for c in wb["الموظفون"][1]} | {re.sub(r"\s*\*$", "", c.value) for c in wb["المناديب"][1]}
    for f in FIELDS:
        assert f.xlsx in heads, f.xlsx


def test_international_emergency_phone():
    f = BY_KEY["emergency_phone"]
    assert clean(f, "+91 98765 43210") == ("+919876543210", None)
    assert clean(f, "0091 98765 43210") == ("+919876543210", None)
    assert clean(f, "+880 1712-345678") == ("+8801712345678", None)
    assert clean(f, "0512345678") == ("+966512345678", None)
    assert clean(f, "123")[1] and clean(f, "abc")[1]


def test_only_english_name_required():
    assert BY_KEY["name_en"].required and not BY_KEY["name_ar"].required
    from app.fields import display_name
    assert display_name({"name_en": "Ravi Kumar"}) == "Ravi Kumar"
    assert display_name({"name_ar": "أحمد", "name_en": "Ahmed"}) == "أحمد"
    assert display_name({}, "E1") == "E1"


def test_delivery_apps_are_the_six_company_apps():
    assert OPTIONS["app"] == ["كيتا فود", "كيتا ماركت", "نينجا", "تويو", "هنقرستيشن", "مرسول"]


def test_translations_complete():
    """Every label/option/message the portal shows exists in all five languages (English at minimum for options)."""
    from app import i18n
    assert i18n._selfcheck() == []
    for key, (ar, table) in i18n.UI.items():
        assert set(table) == {"en", "ur", "bn", "hi"}, key
    for key, table in i18n.SECTION_TITLES.items():
        assert set(table) == {"en", "ur", "bn", "hi"}, key
    for ar, table in i18n.MESSAGES.items():
        assert set(table) == {"en", "ur", "bn", "hi"}, ar
    for lang in i18n.LANGS:
        assert "{name}" in i18n.WA_MESSAGE[lang] and "{link}" in i18n.WA_MESSAGE[lang]
    # every field is labelled in every language; no label is silently the Arabic fallback
    for f in FIELDS:
        if f.key in i18n.FIELD_LABELS:
            continue
        raise AssertionError(f"field {f.key} has no translations")
    for lang in ("en", "ur", "bn", "hi"):
        tr = i18n.Tr(lang)
        for f in FIELDS:
            assert tr.label(f) != f.label, (lang, f.key)
